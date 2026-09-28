import tkinter as tk
from tkinter import ttk, filedialog, messagebox, font as tkfont, colorchooser
import json, os, sys, uuid, shutil, socket, getpass, zipfile, math, re, threading, urllib.request, urllib.error, urllib.parse, subprocess, hashlib, platform, time
from pathlib import Path
from datetime import datetime, date, timedelta

APP_NAME='Ramavtalade tidsplaner'; VERSION='3.3.1'; DATEFMT='%Y-%m-%d'
SUPABASE_URL='https://ipldltuoqstsplnvuijz.supabase.co'
SUPABASE_KEY='sb_publishable_ehE_J4RXZP2G6Y-AYyoCIA_87RGGBLj'
if sys.platform == 'darwin':
    APPDATA=Path.home()/'Library'/'Application Support'/APP_NAME
    LOCAL=Path.home()/'Library'/'Caches'/APP_NAME
else:
    APPDATA=Path(os.getenv('APPDATA', Path.home()))/APP_NAME
    LOCAL=Path(os.getenv('LOCALAPPDATA', Path.home()))/APP_NAME
CFG=APPDATA/'config.json'; BACKUP=LOCAL/'backup'; LICENSE_FILE=APPDATA/'license.json'
STATUSES=['Ej startad','Pågår','Pausad','Klar']
PROJECT_COLOR='#2563eb'
DEFAULT_ACTIVITY_COLOR='#16a34a'

def parse_date(s): return datetime.strptime(s,DATEFMT).date()
def now(): return datetime.now().isoformat(timespec='seconds')
def time_progress(start_s, end_s, today=None):
    """Calendar-time progress, 0-100, based only on today's date."""
    try:
        start, end = parse_date(start_s), parse_date(end_s)
        today = today or date.today()
        if today < start: return 0
        if today >= end: return 100
        span = max(1, (end-start).days)
        return max(0, min(100, round(((today-start).days/span)*100)))
    except Exception:
        return 0
def safe_read(path, default=None):
    try: return json.loads(path.read_text(encoding='utf-8'))
    except: return default

def open_native(path):
    """Open a file/folder with the platform's default application."""
    path=str(path)
    try:
        if sys.platform == 'darwin': subprocess.Popen(['open', path])
        elif os.name == 'nt': os.startfile(path)
        else: subprocess.Popen(['xdg-open', path])
        return True
    except Exception:
        return False

def extract_json(text):
    text=(text or '').strip()
    if text.startswith('```'):
        text=re.sub(r'^```(?:json)?\s*','',text,flags=re.I)
        text=re.sub(r'\s*```$','',text)
    try: return json.loads(text)
    except Exception: pass
    a=text.find('{'); b=text.rfind('}')
    if a>=0 and b>a: return json.loads(text[a:b+1])
    raise ValueError('AI-svaret innehöll ingen giltig JSON.')

def atomic_json(path, data):
    path.parent.mkdir(parents=True,exist_ok=True); tmp=path.with_suffix(path.suffix+'.tmp')
    tmp.write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8'); os.replace(tmp,path)


def _license_state(): return safe_read(LICENSE_FILE,{}) or {}
def _save_license_state(d): LICENSE_FILE.parent.mkdir(parents=True,exist_ok=True); atomic_json(LICENSE_FILE,d)
def _apply_license_profile(res):
    state=_license_state(); state.update({k:res.get(k) for k in ('license_id','customer','expires_at','role','display_name','admin_email','onedrive_folder_name','onedrive_share_url','ai_provider','ai_model','email')}); state['last_ok']=now(); _save_license_state(state)

def account_login(parent=None, force_new=False):
    import rtp_auth
    if force_new: rtp_auth.clear_session()
    if not force_new:
        try:
            t=rtp_auth.token()
            if t:
                res=rtp_auth.activate_device(t)
                if isinstance(res,dict) and res.get('ok'): _apply_license_profile(res); return True
        except Exception: pass
    dlg=tk.Toplevel(parent) if parent else tk.Tk(); dlg.title('Logga in'); dlg.geometry('500x355'); dlg.resizable(False,False)
    if parent and str(parent.state())!='withdrawn': dlg.transient(parent)
    dlg.grab_set(); dlg.deiconify(); dlg.lift()
    f=ttk.Frame(dlg,padding=26); f.pack(fill='both',expand=True)
    ttk.Label(f,text='Ramavtalade tidsplaner',font=(('SF Pro Text' if sys.platform == 'darwin' else 'Segoe UI'),19,'bold')).pack(anchor='w')
    ttk.Label(f,text='Logga in med ditt arbetskonto.',foreground='#64748b').pack(anchor='w',pady=(4,18))
    ttk.Label(f,text='E-post').pack(anchor='w'); email=tk.StringVar(value=rtp_auth.load_session().get('email','')); e=ttk.Entry(f,textvariable=email,width=48); e.pack(fill='x',pady=(3,10))
    ttk.Label(f,text='Lösenord').pack(anchor='w'); pw=tk.StringVar(); pe=ttk.Entry(f,textvariable=pw,show='•'); pe.pack(fill='x',pady=(3,8))
    remember=tk.BooleanVar(value=True); ttk.Checkbutton(f,text='Kom ihåg mig',variable=remember).pack(anchor='w')
    status=ttk.Label(f,text=''); status.pack(anchor='w',pady=(8,2)); result={'ok':False}
    def login():
        if not email.get().strip() or not pw.get(): return
        status.config(text='Loggar in…'); dlg.update_idletasks()
        try:
            sess=rtp_auth.sign_in(email.get(),pw.get(),remember.get()); res=rtp_auth.activate_device(sess['access_token'])
            if not res.get('ok'): raise RuntimeError(res.get('message','Kontot saknar en aktiv företagslicens.'))
            _apply_license_profile(res); result['ok']=True; dlg.destroy()
        except Exception as ex: status.config(text=''); messagebox.showerror('Kunde inte logga in',str(ex),parent=dlg)
    def create():
        em=email.get().strip(); password=pw.get()
        if not em or len(password)<6: messagebox.showinfo('Skapa konto','Fyll i e-post och ett lösenord med minst 6 tecken.',parent=dlg); return
        try:
            x=rtp_auth.sign_up(em,password)
            if x.get('access_token'):
                res=rtp_auth.activate_device(x['access_token'])
                if not res.get('ok'): raise RuntimeError(res.get('message','Din e-post är inte tillagd i något företag ännu.'))
                if remember.get(): rtp_auth.save_session({'access_token':x['access_token'],'refresh_token':x['refresh_token'],'email':em})
                _apply_license_profile(res); result['ok']=True; dlg.destroy()
            else: messagebox.showinfo('Kontrollera e-posten','Kontot är skapat. Bekräfta e-postadressen och logga sedan in.',parent=dlg)
        except Exception as ex: messagebox.showerror('Kunde inte skapa konto',str(ex),parent=dlg)
    def forgot():
        if not email.get().strip(): messagebox.showinfo('Glömt lösenord','Fyll i din e-postadress först.',parent=dlg); return
        try: rtp_auth.recover(email.get()); messagebox.showinfo('Återställning skickad','Kontrollera din e-post för återställningslänken.',parent=dlg)
        except Exception as ex: messagebox.showerror('Fel',str(ex),parent=dlg)
    buttons=ttk.Frame(f); buttons.pack(fill='x',pady=(6,0)); ttk.Button(buttons,text='Skapa konto',command=create).pack(side='left'); ttk.Button(buttons,text='Glömt lösenord',command=forgot).pack(side='left',padx=6); ttk.Button(buttons,text='Logga in',command=login).pack(side='right')
    pe.bind('<Return>',lambda e:login()); (e if not email.get() else pe).focus_set(); dlg.protocol('WM_DELETE_WINDOW',dlg.destroy)
    if parent: parent.wait_window(dlg)
    else: dlg.mainloop()
    return result['ok']

class ItemDialog(tk.Toplevel):
    def __init__(self,parent,title,data=None,is_project=False,people=None,available_trades=None):
        super().__init__(parent); self.result=None; self.title(title); self.resizable(False,False); self.transient(parent); self.grab_set()
        d=data or {}; people=people or []; available_trades=available_trades or []
        self.vars={k:tk.StringVar(value=str(d.get(k,v))) for k,v in {
            'name':'','start':date.today().isoformat(),'end':(date.today()+timedelta(days=14)).isoformat(),
            'status':'Ej startad','owner':'','note':'','color':d.get('color',DEFAULT_ACTIVITY_COLOR),
            'order_no':d.get('order_no',''),'project_no':d.get('project_no',''),'address':d.get('address',''),'handler':d.get('handler',''),
            'bar_text':d.get('bar_text',''),'text_position':d.get('text_position','Automatisk')}.items()}
        frm=ttk.Frame(self,padding=18); frm.pack(fill='both',expand=True)
        fields=[('Namn','name'),('Startdatum (ÅÅÅÅ-MM-DD)','start'),('Slutdatum (ÅÅÅÅ-MM-DD)','end')]
        for r,(lab,key) in enumerate(fields):
            ttk.Label(frm,text=lab).grid(row=r,column=0,sticky='w',pady=5); ttk.Entry(frm,textvariable=self.vars[key],width=42).grid(row=r,column=1,pady=5)
        r=3
        if is_project:
            for lab,key in [('Ordernr','order_no'),('Projektnr','project_no'),('Adress','address'),('Handläggare','handler')]:
                ttk.Label(frm,text=lab).grid(row=r,column=0,sticky='w',pady=5); ttk.Entry(frm,textvariable=self.vars[key],width=42).grid(row=r,column=1,pady=5); r+=1
            ttk.Label(frm,text='Yrkesgrupper').grid(row=r,column=0,sticky='nw',pady=5)
            tradebox=ttk.Frame(frm); tradebox.grid(row=r,column=1,sticky='ew',pady=5)
            self.trade_list=tk.Listbox(tradebox,selectmode='multiple',height=min(7,max(3,len(available_trades))),exportselection=False,width=39)
            self.trade_list.pack(side='left',fill='x',expand=True)
            sb=ttk.Scrollbar(tradebox,orient='vertical',command=self.trade_list.yview); sb.pack(side='right',fill='y'); self.trade_list.configure(yscrollcommand=sb.set)
            selected=set(d.get('trades',[]))
            for i,t in enumerate(available_trades):
                self.trade_list.insert('end',t)
                if t in selected:self.trade_list.selection_set(i)
            r+=1
        ttk.Label(frm,text='Status').grid(row=r,column=0,sticky='w',pady=5); ttk.Combobox(frm,textvariable=self.vars['status'],values=STATUSES,state='readonly',width=39).grid(row=r,column=1,pady=5)
        r+=1; ttk.Label(frm,text='Tidsmässigt klart').grid(row=r,column=0,sticky='w',pady=5); ttk.Label(frm,text=f"{time_progress(self.vars['start'].get(), self.vars['end'].get())}% (automatiskt från datum)").grid(row=r,column=1,sticky='w',pady=5)
        r+=1; ttk.Label(frm,text='Ansvarig').grid(row=r,column=0,sticky='w',pady=5); ttk.Combobox(frm,textvariable=self.vars['owner'],values=people,width=39).grid(row=r,column=1,pady=5)
        r+=1
        if not is_project:
            ttk.Label(frm,text='Stapelfärg').grid(row=r,column=0,sticky='w',pady=5)
            cf=ttk.Frame(frm); cf.grid(row=r,column=1,sticky='w',pady=5)
            self.color_preview=tk.Label(cf,width=4,bg=self.vars['color'].get(),relief='solid',bd=1); self.color_preview.pack(side='left',padx=(0,8))
            def choose_color():
                c=colorchooser.askcolor(color=self.vars['color'].get(),parent=self)[1]
                if c:self.vars['color'].set(c); self.color_preview.configure(bg=c)
            ttk.Button(cf,text='Välj färg…',command=choose_color).pack(side='left')
            r+=1
        ttk.Label(frm,text='Text på/över stapel').grid(row=r,column=0,sticky='w',pady=5); ttk.Entry(frm,textvariable=self.vars['bar_text'],width=42).grid(row=r,column=1,pady=5); r+=1
        ttk.Label(frm,text='Textplacering').grid(row=r,column=0,sticky='w',pady=5); ttk.Combobox(frm,textvariable=self.vars['text_position'],values=['Automatisk','I stapeln','Ovanför','Ingen'],state='readonly',width=39).grid(row=r,column=1,pady=5); r+=1
        ttk.Label(frm,text='Anteckning').grid(row=r,column=0,sticky='nw',pady=5); self.note=tk.Text(frm,width=42,height=5); self.note.grid(row=r,column=1,pady=5); self.note.insert('1.0',d.get('note',''))
        btn=ttk.Frame(frm); btn.grid(row=r+1,column=0,columnspan=2,sticky='e',pady=(14,0)); ttk.Button(btn,text='Avbryt',command=self.destroy).pack(side='right',padx=4,pady=4); ttk.Button(btn,text='Spara',command=self.ok).pack(side='right',padx=4,pady=4)
        self.bind('<Escape>',lambda e:self.destroy()); self.bind('<Return>',lambda e:self.ok()); self.after(100,lambda:self.focus_force())
    def ok(self):
        try:
            s=parse_date(self.vars['start'].get()); e=parse_date(self.vars['end'].get()); assert e>=s
            if not self.vars['name'].get().strip(): raise ValueError
        except: messagebox.showerror('Kontrollera uppgifterna','Ange namn och giltiga datum.',parent=self); return
        self.result={k:v.get().strip() for k,v in self.vars.items()}; self.result['trades']=[self.trade_list.get(i) for i in self.trade_list.curselection()] if hasattr(self,'trade_list') else []; self.result['progress']=time_progress(self.result['start'],self.result['end']); self.result['note']=self.note.get('1.0','end').strip(); self.destroy()

class App(tk.Tk):
    def __init__(self):
        super().__init__(); self.title(f'{APP_NAME}  •  v{VERSION}'); self.geometry('1580x900'); self.minsize(1180,700)
        # Configure Tk named fonts safely; avoid Tcl parsing of multi-word family strings.
        for _name in ('TkDefaultFont','TkTextFont','TkMenuFont','TkHeadingFont','TkCaptionFont','TkSmallCaptionFont','TkIconFont','TkTooltipFont'):
            try: tkfont.nametofont(_name).configure(family=('SF Pro Text' if sys.platform == 'darwin' else 'Segoe UI'), size=9)
            except tk.TclError: pass
        self.cfg=safe_read(CFG,{}) or {}; self.license_profile=_license_state(); self.workspace_id=str(self.license_profile.get('license_id') or '').strip(); self.user=self.license_profile.get('display_name') or self.cfg.get('user') or getpass.getuser()
        # API-nycklar är lokala men strikt isolerade per företags-/workspace-ID.
        # Migrera en äldre global nyckel endast till det företag som är aktivt vid uppgraderingen.
        ai_keys=dict(self.cfg.get('ai_keys') or {})
        legacy_key=str(self.cfg.get('ai_key') or '').strip()
        if self.workspace_id and legacy_key and self.workspace_id not in ai_keys:
            ai_keys[self.workspace_id]=legacy_key
        self.cfg['ai_keys']=ai_keys
        self.cfg['ai_key']=str(ai_keys.get(self.workspace_id,'')).strip() if self.workspace_id else ''
        # OneDrive-path is scoped per company/license. Never reuse a global path from another company.
        self.cfg.setdefault('workspace_sync', {})
        saved_sync=(self.cfg.get('workspace_sync') or {}).get(self.workspace_id, '') if self.workspace_id else ''
        self.sync=Path(saved_sync) if saved_sync else None
        self.cfg['ai_provider']=self.license_profile.get('ai_provider') or self.cfg.get('ai_provider','Gemini'); self.cfg['ai_model']=self.license_profile.get('ai_model') or self.cfg.get('ai_model','gemini-3.5-flash')
        self.projects=[]; self.dirty=set(); self.mtimes={}; self.zoom=18; self.drag=None; self.collapsed=set(self.cfg.get('collapsed',[])); self.logo=Path(self.cfg['logo']) if self.cfg.get('logo') else None; self.filter_text=tk.StringVar(); self.filter_owner=tk.StringVar(value='Alla'); self.filter_status=tk.StringVar(value='Alla'); self.autosave=tk.BooleanVar(value=True)
        self.style_ui(); self.build_ui(); self.ensure_sync(); self.load_all(); self.after(4000,self.poll)
    def style_ui(self):
        s=ttk.Style(self)
        try: s.theme_use('clam')
        except: pass
        self.colors={'nav':'#0f1d31','nav2':'#162a45','bg':'#f4f7fb','card':'#ffffff','text':'#14213d','muted':'#64748b','blue':'#2563eb','line':'#dce5ef','green':'#16a34a','red':'#dc2626'}
        self.configure(bg=self.colors['bg'])
        self.font_title=tkfont.Font(root=self,family=('SF Pro Text' if sys.platform == 'darwin' else 'Segoe UI'),size=20,weight='bold')
        self.font_heading=tkfont.Font(root=self,family=('SF Pro Text' if sys.platform == 'darwin' else 'Segoe UI'),size=9,weight='bold')
        self.font_canvas_bold=tkfont.Font(root=self,family=('SF Pro Text' if sys.platform == 'darwin' else 'Segoe UI'),size=8,weight='bold')
        self.font_canvas=tkfont.Font(root=self,family=('SF Pro Text' if sys.platform == 'darwin' else 'Segoe UI'),size=8)
        self.font_month=tkfont.Font(root=self,family=('SF Pro Text' if sys.platform == 'darwin' else 'Segoe UI'),size=9,weight='bold')
        s.configure('.',font=(('SF Pro Text' if sys.platform == 'darwin' else 'Segoe UI'),9),background=self.colors['bg'],foreground=self.colors['text'])
        s.configure('Main.TFrame',background=self.colors['bg'])
        s.configure('Card.TFrame',background=self.colors['card'])
        s.configure('Title.TLabel',font=self.font_title,background=self.colors['bg'],foreground=self.colors['text'])
        s.configure('Sub.TLabel',background=self.colors['bg'],foreground=self.colors['muted'])
        s.configure('CardTitle.TLabel',background=self.colors['card'],foreground=self.colors['muted'],font=(('SF Pro Text' if sys.platform == 'darwin' else 'Segoe UI'),9))
        s.configure('CardValue.TLabel',background=self.colors['card'],foreground=self.colors['text'],font=(('SF Pro Text' if sys.platform == 'darwin' else 'Segoe UI'),18,'bold'))
        s.configure('Treeview',rowheight=34,background='#ffffff',fieldbackground='#ffffff',foreground=self.colors['text'],borderwidth=0)
        s.configure('Treeview.Heading',font=self.font_heading,background='#f8fafc',foreground='#475569',relief='flat',padding=(6,8))
        s.map('Treeview',background=[('selected','#dbeafe')],foreground=[('selected','#0f172a')])
        s.configure('Primary.TButton',font=(('SF Pro Text' if sys.platform == 'darwin' else 'Segoe UI'),9,'bold'),padding=(12,8),background=self.colors['blue'],foreground='white',borderwidth=0)
        s.map('Primary.TButton',background=[('active','#1d4ed8')])
        s.configure('Tool.TButton',padding=(9,7),background='#ffffff',foreground=self.colors['text'],borderwidth=1)
        s.map('Tool.TButton',background=[('active','#eef4ff')])
    def save_cfg(self):
        CFG.parent.mkdir(parents=True,exist_ok=True)
        ws=dict(self.cfg.get('workspace_sync') or {})
        if self.workspace_id and self.sync: ws[self.workspace_id]=str(self.sync)
        ai_keys=dict(self.cfg.get('ai_keys') or {})
        if self.workspace_id:
            current_key=str(self.cfg.get('ai_key') or '').strip()
            if current_key: ai_keys[self.workspace_id]=current_key
            else: ai_keys.pop(self.workspace_id,None)
        self.cfg['ai_keys']=ai_keys
        atomic_json(CFG,{'workspace_sync':ws, 'user':self.user, 'logo':str(self.logo) if self.logo else '', 'collapsed':sorted(self.collapsed), 'ai_provider':self.cfg.get('ai_provider','Gemini'), 'ai_keys':ai_keys, 'ai_model':self.cfg.get('ai_model','gemini-3.5-flash')})
    def ensure_sync(self):
        # v3.1: företagets licens styr mappnamnet. Sök automatiskt i lokalt synkade OneDrive-rötter.
        wanted=(self.license_profile.get('onedrive_folder_name') or 'Ramavtalade tidsplaner').strip()
        if not self.sync or not self.sync.exists():
            roots=[]
            if sys.platform == 'darwin':
                cloud=Path.home()/'Library'/'CloudStorage'
                if cloud.exists():
                    roots.extend(x for x in cloud.iterdir() if x.is_dir() and 'onedrive' in x.name.lower())
            else:
                for k,v in os.environ.items():
                    if k.lower().startswith('onedrive') and v and Path(v).exists(): roots.append(Path(v))
            candidates=[]
            for root in roots:
                candidates += [root/wanted, root/'Shared'/wanted, root/'Delat'/wanted]
            found=next((c for c in candidates if c.exists()),None)
            if found:
                self.sync=found/'RamavtaladeTidsplanerData'
            else:
                share=self.license_profile.get('onedrive_share_url') or ''
                msg=f'Företagets OneDrive-mapp "{wanted}" är inte synkad på den här datorn.'
                if share: msg+='\n\nÖppna företagets delningslänk och lägg till/synka mappen i OneDrive. Därefter väljer du den lokala mappen här.'
                messagebox.showinfo('Anslut företagets OneDrive',msg,parent=self)
                p=filedialog.askdirectory(title=f'Välj den synkade OneDrive-mappen: {wanted}')
                self.sync=(Path(p)/'RamavtaladeTidsplanerData') if p else (LOCAL/'offline-data')
            self.sync.mkdir(parents=True,exist_ok=True); self.save_cfg()
        (self.sync/'projects').mkdir(parents=True,exist_ok=True); (self.sync/'archive').mkdir(exist_ok=True); (self.sync/'mallar').mkdir(exist_ok=True); (self.sync/'documents').mkdir(exist_ok=True)
    @property
    def pdir(self): return self.sync/'projects'
    @property
    def template_dir(self): return self.sync/'mallar'
    @property
    def documents_dir(self): return self.sync/'documents'
    def available_trades(self):
        self.template_dir.mkdir(parents=True,exist_ok=True)
        return sorted([x.name for x in self.template_dir.iterdir() if x.is_dir()], key=str.lower)
    def project_doc_dir(self,p):
        d=self.documents_dir/(p.get('project_no') or p['id'])
        d.mkdir(parents=True,exist_ok=True)
        return d

    def build_ui(self):
        # Modern shell: dark navigation + light workspace. Functionality remains unchanged.
        shell=tk.Frame(self,bg=self.colors['bg']); shell.pack(fill='both',expand=True)
        nav=tk.Frame(shell,bg=self.colors['nav'],width=205); nav.pack(side='left',fill='y'); nav.pack_propagate(False)
        main=ttk.Frame(shell,style='Main.TFrame'); main.pack(side='left',fill='both',expand=True)
        brand=tk.Frame(nav,bg=self.colors['nav']); brand.pack(fill='x',padx=18,pady=(22,24))
        tk.Label(brand,text='◆',bg=self.colors['nav'],fg='#60a5fa',font=(('SF Pro Text' if sys.platform == 'darwin' else 'Segoe UI'),20,'bold')).pack(side='left')
        tk.Label(brand,text='Ramavtalade\ntidsplaner',justify='left',bg=self.colors['nav'],fg='white',font=(('SF Pro Text' if sys.platform == 'darwin' else 'Segoe UI'),11,'bold')).pack(side='left',padx=9)
        def navbtn(text,cmd,active=False):
            b=tk.Button(nav,text=text,command=cmd,anchor='w',bd=0,relief='flat',cursor='hand2',padx=18,pady=10,
                        bg=('#2563eb' if active else self.colors['nav']),fg='white' if active else '#cbd5e1',
                        activebackground='#1e40af' if active else self.colors['nav2'],activeforeground='white',font=(('SF Pro Text' if sys.platform == 'darwin' else 'Segoe UI'),9,'bold' if active else 'normal'))
            b.pack(fill='x',padx=10,pady=2); return b
        navbtn('⌂   Översikt',self.refresh)
        navbtn('▣   Projekt',self.refresh,True)
        navbtn('▤   Tidsplan / Gantt',lambda: self.canvas.focus_set())
        navbtn('✓   Aktiviteter',lambda: self.tree.focus_set())
        navbtn('▧   Dokument',self.document_manager)
        navbtn('✦   AI-assistent',self.ai_assistant)
        tk.Frame(nav,bg='#24364f',height=1).pack(fill='x',padx=18,pady=16)
        navbtn('⚙   Inställningar',self.settings)
        tk.Label(nav,text=f'v{VERSION}',bg=self.colors['nav'],fg='#64748b',font=(('SF Pro Text' if sys.platform == 'darwin' else 'Segoe UI'),8)).pack(side='bottom',anchor='w',padx=22,pady=18)

        head=ttk.Frame(main,style='Main.TFrame',padding=(24,18,24,8)); head.pack(fill='x')
        lefthead=ttk.Frame(head,style='Main.TFrame'); lefthead.pack(side='left')
        ttk.Label(lefthead,text='Projekt',style='Title.TLabel').pack(anchor='w')
        ttk.Label(lefthead,text='Alla projekt samlade med tidslinje, status och framdrift.',style='Sub.TLabel').pack(anchor='w',pady=(2,0))
        self.sync_lbl=ttk.Label(head,text='',style='Sub.TLabel'); self.sync_lbl.pack(side='right',pady=8)

        cards=ttk.Frame(main,style='Main.TFrame',padding=(24,4,24,8)); cards.pack(fill='x')
        self.stat_labels={}
        for key,title in [('active','Aktiva projekt'),('upcoming','Kommande projekt'),('delayed','Behöver uppmärksamhet'),('progress','Genomsnittlig framdrift')]:
            card=ttk.Frame(cards,style='Card.TFrame',padding=(16,12)); card.pack(side='left',fill='x',expand=True,padx=(0,10))
            ttk.Label(card,text=title,style='CardTitle.TLabel').pack(anchor='w')
            lab=ttk.Label(card,text='0',style='CardValue.TLabel'); lab.pack(anchor='w',pady=(3,0)); self.stat_labels[key]=lab

        tools=ttk.Frame(main,style='Main.TFrame',padding=(24,6,24,6)); tools.pack(fill='x')
        ttk.Button(tools,text='＋ Nytt projekt',command=self.add_project,style='Primary.TButton').pack(side='left',padx=(0,6))
        for t,c in [('＋ Aktivitet',self.add_activity),('Redigera',self.edit_selected),('Duplicera',self.duplicate_selected),('Ta bort',self.delete_selected),('Spara',self.save_all),('↻ Synka',self.load_all)]: ttk.Button(tools,text=t,command=c,style='Tool.TButton').pack(side='left',padx=3)
        ttk.Button(tools,text='PDF A3',command=self.export_pdf,style='Tool.TButton').pack(side='right',padx=3)
        ttk.Button(tools,text='−',width=3,command=lambda:self.set_zoom(-3),style='Tool.TButton').pack(side='right'); ttk.Button(tools,text='+',width=3,command=lambda:self.set_zoom(3),style='Tool.TButton').pack(side='right',padx=(3,0)); ttk.Label(tools,text='Zoom',style='Sub.TLabel').pack(side='right',padx=5)

        filt=ttk.Frame(main,style='Main.TFrame',padding=(24,5,24,9)); filt.pack(fill='x')
        ttk.Label(filt,text='🔍',style='Sub.TLabel').pack(side='left'); self.search_entry=ttk.Entry(filt,textvariable=self.filter_text,width=40); self.search_entry.pack(side='left',padx=(6,16)); self.search_entry.insert(0,''); self.search_entry.bind('<KeyRelease>',lambda e:self.refresh())
        self.search_entry.bind('<Escape>',lambda e:self.clear_search())
        self.bind_all('<Control-f>',lambda e:self.focus_search())
        ttk.Label(filt,text='Projekt / adress / ordernr / projektnr',style='Sub.TLabel').pack(side='left',padx=(0,18))
        ttk.Label(filt,text='Ansvarig',style='Sub.TLabel').pack(side='left'); self.owner_cb=ttk.Combobox(filt,textvariable=self.filter_owner,state='readonly',width=16); self.owner_cb.pack(side='left',padx=(6,16)); self.owner_cb.bind('<<ComboboxSelected>>',lambda e:self.refresh())
        ttk.Label(filt,text='Status',style='Sub.TLabel').pack(side='left'); st=ttk.Combobox(filt,textvariable=self.filter_status,values=['Alla']+STATUSES,state='readonly',width=13); st.pack(side='left',padx=6); st.bind('<<ComboboxSelected>>',lambda e:self.refresh())
        self.summary=ttk.Label(filt,text='',style='Sub.TLabel'); self.summary.pack(side='right')

        content=ttk.Frame(main,style='Card.TFrame',padding=1); content.pack(fill='both',expand=True,padx=24,pady=(0,10))
        pan=ttk.Panedwindow(content,orient='horizontal'); pan.pack(fill='both',expand=True); left=ttk.Frame(pan,style='Card.TFrame'); right=ttk.Frame(pan,style='Card.TFrame'); pan.add(left,weight=0); pan.add(right,weight=1)
        cols=('order_no','project_no','handler','owner','status','progress','start','end'); self.tree=ttk.Treeview(left,columns=cols,show='tree headings',selectmode='browse')
        headers={'#0':'Projekt / aktivitet','order_no':'Ordernr','project_no':'Projektnr','handler':'Handläggare','owner':'Ansvarig','status':'Status','progress':'Tid %','start':'Start','end':'Slut'}
        widths={'#0':205,'order_no':76,'project_no':82,'handler':100,'owner':100,'status':78,'progress':55,'start':78,'end':78}
        for col in ['#0']+list(cols): self.tree.heading(col,text=headers[col]); self.tree.column(col,width=widths[col],anchor='w')
        self.tree.pack(fill='both',expand=True); self.tree.bind('<Double-1>',lambda e:self.edit_selected()); self.tree.bind('<<TreeviewSelect>>',lambda e:self.draw()); self.tree.bind('<<TreeviewOpen>>',self.tree_toggle); self.tree.bind('<<TreeviewClose>>',self.tree_toggle)
        self.canvas=tk.Canvas(right,bg='#ffffff',highlightthickness=0); hx=ttk.Scrollbar(right,orient='horizontal',command=self.canvas.xview); vy=ttk.Scrollbar(right,orient='vertical',command=self.canvas.yview); self.canvas.configure(xscrollcommand=hx.set,yscrollcommand=vy.set); self.canvas.grid(row=0,column=0,sticky='nsew'); vy.grid(row=0,column=1,sticky='ns'); hx.grid(row=1,column=0,sticky='ew'); right.rowconfigure(0,weight=1); right.columnconfigure(0,weight=1)
        self.canvas.bind('<ButtonPress-1>',self.drag_start); self.canvas.bind('<B1-Motion>',self.drag_move); self.canvas.bind('<ButtonRelease-1>',self.drag_end); self.canvas.bind('<MouseWheel>',lambda e:self.canvas.yview_scroll(int(-e.delta/120),'units'))
        foot=ttk.Frame(main,style='Main.TFrame',padding=(24,2,24,8)); foot.pack(fill='x'); self.status=ttk.Label(foot,text='',style='Sub.TLabel'); self.status.pack(side='left'); ttk.Label(foot,text='Dra en aktivitetsrad i tidslinjen för att flytta den i kalendern.',style='Sub.TLabel').pack(side='right')
    def focus_search(self):
        if hasattr(self,'search_entry'):
            self.search_entry.focus_set(); self.search_entry.selection_range(0,'end')
        return 'break'
    def clear_search(self):
        self.filter_text.set(''); self.refresh(); self.focus_search(); return 'break'
    def people(self): return sorted({x.get('owner','') for p in self.projects for x in [p]+p.get('activities',[]) if x.get('owner')})
    def filtered(self):
        q=self.filter_text.get().casefold().strip(); own=self.filter_owner.get(); st=self.filter_status.get(); out=[]
        for p in sorted(self.projects,key=lambda x:x.get('start','')):
            # Snabbsökningen avser projektidentitet, inte AI eller dokumentinnehåll.
            hay=' '.join(str(p.get(k,'') or '') for k in ('name','address','order_no','project_no')).casefold(); acts=p.get('activities',[])
            if q and q not in hay: continue
            if own!='Alla' and p.get('owner')!=own and not any(a.get('owner')==own for a in acts): continue
            if st!='Alla' and p.get('status')!=st and not any(a.get('status')==st for a in acts): continue
            out.append(p)
        return out
    def load_all(self):
        if self.dirty:
            if not messagebox.askyesno('Osparade ändringar','Det finns osparade ändringar. Hämta från OneDrive och kasta dem?'): return
        ps=[]; mt={}
        for f in self.pdir.glob('*.json'):
            d=safe_read(f); 
            if isinstance(d,dict) and d.get('id'): ps.append(d); mt[d['id']]=f.stat().st_mtime
        # migrate old projects.json if present and no split files
        old=self.sync/'projects.json'
        if not ps and old.exists():
            raw=safe_read(old,{}) or {}
            for p in raw.get('projects',[]): self.normalize(p); atomic_json(self.pdir/f"{p['id']}.json",p)
            if raw.get('projects'): old.rename(self.sync/'archive'/f'projects_legacy_{datetime.now():%Y%m%d_%H%M%S}.json'); return self.load_all()
        self.projects=ps; self.mtimes=mt; self.dirty.clear(); self.refresh(); self.status.config(text=f'Synkad {datetime.now():%H:%M:%S}'); self.sync_lbl.config(text=f'● OneDrive-mapp  •  {self.user}')
    def normalize(self,p):
        p.setdefault('id',uuid.uuid4().hex); p.setdefault('status','Ej startad'); p.setdefault('owner',''); p.setdefault('order_no',''); p.setdefault('project_no',''); p.setdefault('address',''); p.setdefault('handler',''); p.setdefault('note',''); p.setdefault('trades',[]); p.setdefault('bar_text',''); p.setdefault('text_position','Automatisk'); p.setdefault('activities',[]); p.setdefault('revision',0); p['color']=PROJECT_COLOR
        for a in p['activities']:
            a.setdefault('id',uuid.uuid4().hex); a.setdefault('status','Ej startad'); a.setdefault('owner',p.get('owner','')); a.setdefault('note',''); a.setdefault('color',DEFAULT_ACTIVITY_COLOR); a.setdefault('bar_text',''); a.setdefault('text_position','Automatisk'); a['progress']=time_progress(a.get('start',''),a.get('end',''))
        p['progress']=time_progress(p.get('start',''),p.get('end',''))
    def save_project(self,p):
        f=self.pdir/f"{p['id']}.json"; remote=f.stat().st_mtime if f.exists() else 0; known=self.mtimes.get(p['id'],0)
        if remote>known+0.001:
            remote_data=safe_read(f,{}) or {}; who=remote_data.get('updated_by','annan användare')
            if not messagebox.askyesno('Samtidig ändring',f"{p['name']} har ändrats av {who} efter att du öppnade den.\n\nVill du ändå ersätta den med din version?"): return False
        BACKUP.mkdir(parents=True,exist_ok=True)
        if f.exists(): shutil.copy2(f,BACKUP/f"{p['id']}_{datetime.now():%Y%m%d_%H%M%S}.json")
        p['updated_at']=now(); p['updated_by']=self.user; p['machine']=socket.gethostname(); p['revision']=int(p.get('revision',0))+1; atomic_json(f,p); self.mtimes[p['id']]=f.stat().st_mtime; return True
    def save_all(self):
        ok=True
        for pid in list(self.dirty):
            p=next((x for x in self.projects if x['id']==pid),None)
            if p and self.save_project(p): self.dirty.discard(pid)
            elif p: ok=False
        self.status.config(text=('Sparat och synkat' if ok else 'Vissa ändringar sparades inte')+f'  •  {datetime.now():%H:%M:%S}'); self.refresh()
    def changed(self,p):
        p['updated_at']=now(); p['updated_by']=self.user; self.dirty.add(p['id']); self.refresh(); self.status.config(text='● Osparade ändringar');
        if self.autosave.get(): self.after(350,lambda:self.save_all() if self.dirty else None)
    def poll(self):
        try:
            if not self.dirty:
                current={f.stem:f.stat().st_mtime for f in self.pdir.glob('*.json')}
                if current!=self.mtimes:self.load_all()
        finally:self.after(4000,self.poll)
    def selected(self):
        sel=self.tree.selection(); iid=sel[0] if sel else None
        for p in self.projects:
            if iid==p['id']: return p,None
            for a in p.get('activities',[]):
                if iid==a['id']: return p,a
        return None,None
    def dialog(self,title,data=None,is_project=False):
        d=ItemDialog(self,title,data,is_project,self.people(),self.available_trades()); self.wait_window(d); return d.result
    def add_project(self):
        x=self.dialog('Nytt projekt',{'status':'Ej startad'},True)
        if not x:return
        p={'id':uuid.uuid4().hex,'activities':[],'revision':0,**x}; p['color']=PROJECT_COLOR; self.projects.append(p); self.sync_project_templates(p,quiet=True); self.changed(p)
    def add_activity(self):
        p,a=self.selected()
        if not p: messagebox.showinfo('Välj projekt','Markera projektet som aktiviteten ska tillhöra.'); return
        x=self.dialog('Ny aktivitet',{'start':p['start'],'end':p['end'],'owner':p.get('owner',''),'status':'Ej startad','color':DEFAULT_ACTIVITY_COLOR})
        if not x:return
        p.setdefault('activities',[]).append({'id':uuid.uuid4().hex,**x}); self.recalc_project(p); self.changed(p)
    def edit_selected(self):
        p,a=self.selected(); obj=a or p
        if not obj:return
        x=self.dialog('Redigera '+('aktivitet' if a else 'projekt'),obj,a is None)
        if not x:return
        obj.update(x); self.recalc_project(p); self.sync_project_templates(p,quiet=True) if a is None else None; self.changed(p)
    def duplicate_selected(self):
        p,a=self.selected()
        if not p:return
        if a:
            b=json.loads(json.dumps(a)); b['id']=uuid.uuid4().hex; b['name']+=' – kopia'; p['activities'].append(b); self.changed(p)
        else:
            q=json.loads(json.dumps(p)); q['id']=uuid.uuid4().hex; q['name']+=' – kopia'; q['revision']=0
            for a2 in q.get('activities',[]): a2['id']=uuid.uuid4().hex
            self.projects.append(q); self.changed(q)
    def delete_selected(self):
        p,a=self.selected()
        if not p:return
        if not messagebox.askyesno('Ta bort',f"Ta bort {'aktiviteten' if a else 'projektet'} permanent?"):return
        if a:p['activities']=[x for x in p['activities'] if x['id']!=a['id']]; self.recalc_project(p); self.changed(p)
        else:
            f=self.pdir/f"{p['id']}.json"; 
            if f.exists(): shutil.move(str(f),str(self.sync/'archive'/f"{p['id']}_{datetime.now():%Y%m%d_%H%M%S}.json"))
            self.projects=[x for x in self.projects if x['id']!=p['id']]; self.mtimes.pop(p['id'],None); self.dirty.discard(p['id']); self.refresh()
    def recalc_project(self,p):
        acts=p.get('activities',[])
        if acts:
            p['start']=min(a['start'] for a in acts); p['end']=max(a['end'] for a in acts); p['progress']=time_progress(p['start'],p['end']);
            for a in acts: a['progress']=time_progress(a['start'],a['end'])
            if all(a.get('status')=='Klar' for a in acts):p['status']='Klar'
            elif any(a.get('status')=='Pågår' for a in acts):p['status']='Pågår'
    def refresh(self):
        self.refresh_tree(); self.draw(); self.owner_cb['values']=['Alla']+self.people()
        acts=sum(len(p.get('activities',[])) for p in self.projects); self.summary.config(text=f"{len(self.projects)} projekt  •  {acts} aktiviteter")
        if hasattr(self,'stat_labels'):
            today=date.today(); active=sum(1 for p in self.projects if p.get('status') in ('Pågår','Pausad')); upcoming=sum(1 for p in self.projects if p.get('status')=='Ej startad' and p.get('start','9999-99-99')>=today.isoformat()); attention=sum(1 for p in self.projects if p.get('status')!='Klar' and p.get('end','9999-99-99')<today.isoformat()); avg=round(sum(time_progress(p.get('start',''),p.get('end','')) for p in self.projects)/len(self.projects)) if self.projects else 0
            self.stat_labels['active'].config(text=str(active)); self.stat_labels['upcoming'].config(text=str(upcoming)); self.stat_labels['delayed'].config(text=str(attention)); self.stat_labels['progress'].config(text=f'{avg}%')
    def refresh_tree(self):
        self.tree.delete(*self.tree.get_children())
        for p in self.filtered():
            self.normalize(p); self.tree.insert('', 'end',iid=p['id'],text=p['name'],values=(p.get('order_no',''),p.get('project_no',''),p.get('handler',''),p['owner'],p['status'],f"{time_progress(p['start'],p['end'])}%",p['start'],p['end']),open=(p['id'] not in self.collapsed))
            for a in sorted(p.get('activities',[]),key=lambda x:x['start']): self.tree.insert(p['id'],'end',iid=a['id'],text='   '+a['name'],values=('','','',a.get('owner',''),a.get('status',''),f"{time_progress(a['start'],a['end'])}%",a['start'],a['end']))
    def bounds(self):
        objs=[x for p in self.filtered() for x in [p]+p.get('activities',[])]; t=date.today()
        if not objs:return t-timedelta(days=14),t+timedelta(days=60)
        return min(parse_date(x['start']) for x in objs)-timedelta(days=7),max(parse_date(x['end']) for x in objs)+timedelta(days=14)
    def rows(self):
        r=[]
        for p in self.filtered():
            r.append((p,p,True))
            if p['id'] not in self.collapsed: r += [(p,a,False) for a in sorted(p.get('activities',[]),key=lambda x:x['start'])]
        return r
    def set_zoom(self,d): self.zoom=max(8,min(48,self.zoom+d)); self.draw()
    def draw(self):
        c=self.canvas;c.delete('all'); start,end=self.bounds(); days=(end-start).days+1; rowh=32; header=64; w=days*self.zoom; rows=self.rows(); self.draw_meta={'start':start,'rowh':rowh,'header':header,'rows':rows}
        d=start
        for i in range(days):
            x=i*self.zoom
            if d.weekday()>=5:c.create_rectangle(x,0,x+self.zoom,header+len(rows)*rowh,fill='#f8fafc',outline='')
            if d.day==1:c.create_rectangle(x,0,min(w,x+31*self.zoom),23,fill='#eef2f7',outline=''); c.create_text(x+5,11,anchor='w',text=d.strftime('%B %Y').capitalize(),font=self.font_month,fill='#334155')
            if d.weekday()==0:c.create_line(x,23,x,header+len(rows)*rowh,fill='#cbd5e1'); c.create_text(x+4,34,anchor='nw',text=f'v{d.isocalendar().week}',font=self.font_canvas_bold,fill='#475569')
            if self.zoom>=16:c.create_text(x+self.zoom/2,54,text=str(d.day),font=self.font_canvas,fill='#64748b')
            d+=timedelta(days=1)
        for r,(p,obj,isproj) in enumerate(rows):
            y=header+r*rowh; c.create_line(0,y,w,y,fill='#edf0f4'); s=parse_date(obj['start']); e=parse_date(obj['end']); x1=(s-start).days*self.zoom+2; x2=((e-start).days+1)*self.zoom-2; col=PROJECT_COLOR if isproj else obj.get('color',DEFAULT_ACTIVITY_COLOR)
            c.create_rectangle(x1,y+6,x2,y+26,fill=col,outline='',tags=('bar',obj['id']))
            prog=time_progress(obj.get('start',''),obj.get('end','')); px=x1+(x2-x1)*prog/100
            if prog:c.create_rectangle(x1,y+21,px,y+26,fill='#0f172a',outline='',tags=('bar',obj['id']))
            label=(obj.get('bar_text') or (obj['name'] if isproj else '')).strip(); pos=obj.get('text_position','Automatisk')
            if label and pos!='Ingen':
                # Automatic: inside when the bar is wide enough, otherwise just above it.
                inside = pos=='I stapeln' or (pos=='Automatisk' and x2-x1>=max(60,7*len(label)+12))
                if inside:
                    c.create_text(x1+6,y+16,anchor='w',text=label,fill='white',font=self.font_canvas_bold if isproj else self.font_canvas,tags=('bar',obj['id']))
                else:
                    c.create_text(x1+3,y+4,anchor='sw',text=label,fill='#334155',font=self.font_canvas_bold if isproj else self.font_canvas,tags=('bar',obj['id']))
        tx=(date.today()-start).days*self.zoom+self.zoom/2
        if 0<=tx<=w:c.create_line(tx,0,tx,header+len(rows)*rowh,fill='#dc2626',width=2); c.create_text(tx+4,27,anchor='nw',text='IDAG',fill='#dc2626',font=self.font_canvas_bold)
        c.configure(scrollregion=(0,0,max(w,900),max(header+len(rows)*rowh+30,550)))
    def drag_start(self,e):
        x=self.canvas.canvasx(e.x); y=self.canvas.canvasy(e.y); m=self.draw_meta; idx=int((y-m['header'])//m['rowh'])
        if 0<=idx<len(m['rows']):
            p,obj,isproj=m['rows'][idx]
            if not isproj:self.drag=(p,obj,x,obj['start'],obj['end'])
    def drag_move(self,e):
        if not self.drag:return
        dx=self.canvas.canvasx(e.x)-self.drag[2]; days=round(dx/self.zoom); p,obj,_,s,e0=self.drag; ns=parse_date(s)+timedelta(days=days); ne=parse_date(e0)+timedelta(days=days); obj['start']=ns.isoformat(); obj['end']=ne.isoformat(); self.draw()
    def drag_end(self,e):
        if not self.drag:return
        p=self.drag[0]; self.drag=None; self.recalc_project(p); self.changed(p)
    def tree_toggle(self,e=None):
        iid=self.tree.focus()
        p=next((x for x in self.projects if x['id']==iid),None)
        if not p:return
        # Tk updates the open flag after the virtual event; defer reading it.
        def apply():
            if self.tree.item(iid,'open'): self.collapsed.discard(iid)
            else: self.collapsed.add(iid)
            self.save_cfg(); self.draw()
        self.after_idle(apply)
    def collapse_all(self):
        self.collapsed={p['id'] for p in self.filtered()}; self.save_cfg(); self.refresh_tree(); self.draw()
    def expand_all(self):
        self.collapsed.clear(); self.save_cfg(); self.refresh_tree(); self.draw()
    def export_pdf(self):
        p,a=self.selected()
        if not p:
            messagebox.showinfo('Välj projekt','Markera huvudprojektet eller en aktivitet i projektet som ska exporteras.'); return
        try:
            from reportlab.pdfgen import canvas as pdfcanvas
            from reportlab.lib.pagesizes import A3, landscape
            from reportlab.lib.colors import HexColor, black, white
            from reportlab.lib.utils import ImageReader
        except Exception:
            messagebox.showerror('PDF-stöd saknas','PDF-biblioteket ReportLab saknas i denna version. Bygg EXE-filen igen med build_exe.bat.'); return
        out=filedialog.asksaveasfilename(defaultextension='.pdf',filetypes=[('PDF','*.pdf')],initialfile=f"{p['name']}_tidsplan.pdf")
        if not out:return
        try:
            page_w,page_h=landscape(A3); c=pdfcanvas.Canvas(out,pagesize=(page_w,page_h)); margin=30
            # Header
            logo_w=0
            if self.logo and self.logo.exists():
                try:
                    ir=ImageReader(str(self.logo)); iw,ih=ir.getSize(); maxw,maxh=125,48; scale=min(maxw/iw,maxh/ih); dw,dh=iw*scale,ih*scale
                    c.drawImage(ir,margin,page_h-margin-dh,width=dw,height=dh,mask='auto',preserveAspectRatio=True); logo_w=dw+18
                except: pass
            c.setFont('Helvetica-Bold',18); c.drawString(margin+logo_w,page_h-margin-16,p['name'])
            c.setFont('Helvetica',8); meta=f"Ordernr: {p.get('order_no') or '-'}    Projektnr: {p.get('project_no') or '-'}    Handläggare: {p.get('handler') or '-'}"
            c.drawString(margin+logo_w,page_h-margin-31,meta)
            c.setFont('Helvetica',8); meta2=f"Ansvarig: {p.get('owner') or '-'}    Status: {p.get('status','-')}    Tidsmässigt klart: {time_progress(p['start'],p['end'])}%    Exporterad: {datetime.now():%Y-%m-%d %H:%M}"
            c.drawString(margin+logo_w,page_h-margin-43,meta2)
            c.setStrokeColor(HexColor('#cbd5e1')); c.line(margin,page_h-margin-66,page_w-margin,page_h-margin-66)
            acts=sorted(p.get('activities',[]),key=lambda x:x['start']); include_acts=p['id'] not in self.collapsed
            rows=[(p,True)]+([(x,False) for x in acts] if include_acts else [])
            start=min(parse_date(x['start']) for x,_ in rows)-timedelta(days=3); end=max(parse_date(x['end']) for x,_ in rows)+timedelta(days=3); days=max(1,(end-start).days+1)
            left=margin+170; right=page_w-margin; top=page_h-margin-90; bottom=margin+30; header_h=44; avail_h=top-bottom-header_h; row_h=min(26,max(14,avail_h/max(1,len(rows))))
            day_w=(right-left)/days
            # timeline header/weekends
            for i in range(days):
                d=start+timedelta(days=i); x=left+i*day_w
                if d.weekday()>=5:
                    c.setFillColor(HexColor('#f8fafc')); c.rect(x,bottom,day_w,top-bottom,stroke=0,fill=1)
                if d.weekday()==0:
                    c.setStrokeColor(HexColor('#cbd5e1')); c.line(x,bottom,x,top); c.setFillColor(HexColor('#475569')); c.setFont('Helvetica-Bold',7); c.drawString(x+2,top-12,f"v{d.isocalendar().week}")
                if day_w>=8:
                    c.setFillColor(HexColor('#64748b')); c.setFont('Helvetica',6); c.drawCentredString(x+day_w/2,top-29,str(d.day))
            # rows
            y=top-header_h
            for obj,isproj in rows:
                c.setStrokeColor(HexColor('#e2e8f0')); c.line(margin,y,right,y)
                c.setFillColor(HexColor('#0f172a')); c.setFont('Helvetica-Bold' if isproj else 'Helvetica',8)
                label=obj['name']; c.drawString(margin+4,y-row_h/2+2,label[:38])
                x1=left+(parse_date(obj['start'])-start).days*day_w; x2=left+((parse_date(obj['end'])-start).days+1)*day_w
                col=PROJECT_COLOR if isproj else obj.get('color',DEFAULT_ACTIVITY_COLOR); c.setFillColor(HexColor(col)); c.roundRect(x1,y-row_h+4,max(2,x2-x1),max(5,row_h-8),3,stroke=0,fill=1)
                bar_label=(obj.get('bar_text') or (obj['name'] if isproj else '')).strip(); pos=obj.get('text_position','Automatisk')
                if bar_label and pos!='Ingen':
                    c.setFont('Helvetica-Bold' if isproj else 'Helvetica',6.5)
                    text_w=c.stringWidth(bar_label,'Helvetica-Bold' if isproj else 'Helvetica',6.5)
                    inside = pos=='I stapeln' or (pos=='Automatisk' and x2-x1>=text_w+10)
                    if inside:
                        c.setFillColor(white); c.drawString(x1+4,y-row_h/2+1,bar_label[:45])
                    else:
                        c.setFillColor(HexColor('#334155')); c.drawString(x1+2,y-5,bar_label[:45])
                y-=row_h
            # today line
            tx=left+(date.today()-start).days*day_w+day_w/2
            if left<=tx<=right:
                c.setStrokeColor(HexColor('#dc2626')); c.setLineWidth(1.5); c.line(tx,bottom,tx,top); c.setFillColor(HexColor('#dc2626')); c.setFont('Helvetica-Bold',7); c.drawString(tx+3,top-42,'IDAG')
            c.setFillColor(HexColor('#64748b')); c.setFont('Helvetica',7); mode='Huvudprojekt + aktiviteter' if include_acts else 'Endast huvudprojekt'; c.drawRightString(right,margin-2,f"A3 liggande - {mode} - Ramavtalade tidsplaner v{VERSION}")
            c.save(); messagebox.showinfo('PDF klar',f'PDF-filen har skapats:\n{out}')
        except Exception as ex: messagebox.showerror('PDF-fel',str(ex))
    def ai_settings_values(self):
        provider=self.cfg.get('ai_provider','Gemini'); key=self.cfg.get('ai_key','').strip(); model=self.cfg.get('ai_model','').strip()
        if not model: model='gemini-3.5-flash' if provider=='Gemini' else 'gpt-5-mini'
        return provider,key,model

    def ai_call(self, prompt):
        provider,key,model=self.ai_settings_values()
        if not key: raise ValueError('Ingen API-nyckel är sparad. Öppna Inställningar och fyll i AI-inställningarna.')
        system=('Du är en projektplaneringsassistent för svensk byggentreprenad. Var konkret och använd ISO-datum YYYY-MM-DD. '
                'Hitta inte på ordernummer, projektnummer eller personer. Användaren fattar alltid beslutet om ändringar.')
        if provider=='OpenAI':
            body={'model':model,'input':[{'role':'system','content':[{'type':'input_text','text':system}]},{'role':'user','content':[{'type':'input_text','text':prompt}]}]}
            req=urllib.request.Request('https://api.openai.com/v1/responses',data=json.dumps(body).encode(),headers={'Authorization':'Bearer '+key,'Content-Type':'application/json'})
            with urllib.request.urlopen(req,timeout=60) as r: data=json.loads(r.read().decode())
            if data.get('output_text'): return data['output_text']
            return '\n'.join(c.get('text','') for o in data.get('output',[]) for c in o.get('content',[]) if c.get('type')=='output_text')
        body={'system_instruction':{'parts':[{'text':system}]},'contents':[{'parts':[{'text':prompt}]}],'generationConfig':{'temperature':0.2}}
        url='https://generativelanguage.googleapis.com/v1beta/models/'+urllib.parse.quote(model,safe='')+':generateContent'
        req=urllib.request.Request(url,data=json.dumps(body).encode(),headers={'x-goog-api-key':key,'Content-Type':'application/json'})
        with urllib.request.urlopen(req,timeout=60) as r: data=json.loads(r.read().decode())
        return ''.join(part.get('text','') for cand in data.get('candidates',[]) for part in cand.get('content',{}).get('parts',[]))

    def project_context(self, p=None):
        # Fresh project facts for AI. Chat memory never replaces current project data.
        if p: return json.dumps(p,ensure_ascii=False,indent=2)
        compact=[]
        for x in self.projects:
            base={k:x.get(k,'') for k in ('id','name','address','order_no','project_no','handler','owner','status','start','end','note')}
            base['progress']=time_progress(x.get('start',''),x.get('end',''))
            base['activities']=[{k:a.get(k,'') for k in ('name','owner','status','start','end','note')} for a in x.get('activities',[])]
            compact.append(base)
        return json.dumps(compact,ensure_ascii=False,indent=2)

    def _chat_file(self):
        email=str(self.license_profile.get('email') or self.user or 'user').strip().casefold()
        uid=hashlib.sha256(email.encode('utf-8')).hexdigest()[:20]
        wid=re.sub(r'[^A-Za-z0-9_.-]+','_',self.workspace_id or 'no_workspace')
        return APPDATA/'chats'/wid/(uid+'.json')

    def load_chat(self):
        d=safe_read(self._chat_file(),{}) or {}; msgs=d.get('messages',[])
        return [m for m in msgs if isinstance(m,dict) and m.get('role') in ('user','assistant') and isinstance(m.get('text'),str)][-40:]

    def save_chat(self, messages):
        atomic_json(self._chat_file(),{'version':1,'workspace_id':self.workspace_id,'user':self.license_profile.get('email') or self.user,'updated_at':now(),'messages':messages[-40:]})

    def chat_prompt(self, question, history):
        ident={'name':self.license_profile.get('display_name') or self.user,'email':self.license_profile.get('email') or '', 'role':self.license_profile.get('role') or '', 'today':date.today().isoformat()}
        transcript='\n'.join(('ANVÄNDARE: ' if m['role']=='user' else 'ASSISTENT: ')+m['text'] for m in history[-16:])
        scope=self.project_context()
        return f'''Du är användarens projektassistent. Förstå följdfrågor utifrån CHATTHISTORIKEN, men använd alltid AKTUELL PROJEKTDATA som källa för projektfakta. Om historiken och projektdata skiljer sig gäller projektdata.\n\nINLOGGAD ANVÄNDARE:\n{json.dumps(ident,ensure_ascii=False)}\n\nCHATTHISTORIK:\n{transcript or "(ingen tidigare historik)"}\n\nAKTUELL PROJEKTDATA:\n{scope}\n\nNY FRÅGA:\n{question}\n\nSvara kort och konkret på svenska. Vid frågor som "vad gör jag idag" ska du prioritera aktiviteter/projekt där ansvarig matchar den inloggade användaren och dagens datum ligger inom start/slut. Om namnmatchning är osäker, säg det. Hitta inte på uppgifter. Ändra aldrig projektdata från chattläget.'''

    def ai_assistant(self):
        win=tk.Toplevel(self); win.title('Projektassistent'); win.geometry('1000x760'); win.transient(self)
        outer=ttk.Frame(win,padding=12); outer.pack(fill='both',expand=True); p,_=self.selected(); provider,key,model=self.ai_settings_values()
        top=ttk.Frame(outer); top.pack(fill='x'); ttk.Label(top,text='Projektassistent',font=(('SF Pro Text' if sys.platform == 'darwin' else 'Segoe UI'),16,'bold')).pack(side='left'); ttk.Label(top,text=f'{provider} • {model}',style='Sub.TLabel').pack(side='right')
        modes=['Chatt','Skapa nytt projekt med AI','Föreslå ändringar i markerat projekt']; mode=tk.StringVar(value=modes[0]); ttk.Combobox(outer,textvariable=mode,values=modes,state='readonly').pack(fill='x',pady=(10,6))
        ttk.Label(outer,text=('Markerat projekt: '+p['name']) if p else 'Chattläge: alla projekt',style='Sub.TLabel').pack(anchor='w',pady=(0,6))
        out=tk.Text(outer,wrap='word',state='disabled',font=(('SF Pro Text' if sys.platform == 'darwin' else 'Segoe UI'),10),background='#ffffff',relief='solid',borderwidth=1); out.pack(fill='both',expand=True,pady=(3,8))
        entry=tk.Text(outer,height=4,wrap='word',font=(('SF Pro Text' if sys.platform == 'darwin' else 'Segoe UI'),10)); entry.pack(fill='x',pady=(0,8))
        buttons=ttk.Frame(outer); buttons.pack(fill='x'); status=ttk.Label(buttons,text=''); status.pack(side='left')
        apply_btn=ttk.Button(buttons,text='Godkänn förslag',state='disabled'); apply_btn.pack(side='right',padx=4); send_btn=ttk.Button(buttons,text='Skicka',style='Primary.TButton'); send_btn.pack(side='right',padx=4)
        history=self.load_chat(); state={'proposal':None}
        def render_chat():
            out.configure(state='normal'); out.delete('1.0','end')
            if not history: out.insert('end','Projektassistent\n\nFråga t.ex. “Vad gör jag idag?”, “Vad händer nästa vecka?” eller “Vilka projekt behöver min uppmärksamhet?”\n')
            else:
                for m in history: out.insert('end',('Du' if m['role']=='user' else 'Assistent')+'\n'+m['text']+'\n\n')
            out.configure(state='disabled'); out.see('end')
        render_chat()
        def new_chat():
            if history and not messagebox.askyesno('Ny chatt','Starta en ny chatt och rensa den sparade chatthistoriken?',parent=win): return
            history.clear(); self.save_chat(history); state['proposal']=None; apply_btn.config(state='disabled'); render_chat(); entry.focus_set()
        ttk.Button(buttons,text='Ny chatt',command=new_chat).pack(side='left',padx=(8,0))
        def setout(t): out.configure(state='normal'); out.delete('1.0','end'); out.insert('1.0',t); out.configure(state='disabled')
        def worker(user_text, selected_mode):
            try:
                if selected_mode=='Chatt':
                    ans=self.ai_call(self.chat_prompt(user_text,history[:-1])); history.append({'role':'assistant','text':ans,'at':now()}); self.save_chat(history)
                    self.after(0,lambda:(render_chat(),status.config(text='Klart'),send_btn.config(state='normal'))); return
                if selected_mode=='Skapa nytt projekt med AI':
                    prompt='''Skapa ett projektutkast från användarens beskrivning. Returnera ENDAST JSON enligt detta schema:\n{"name":"","address":"","start":"YYYY-MM-DD","end":"YYYY-MM-DD","owner":"","handler":"","order_no":"","project_no":"","status":"Ej startad","bar_text":"","activities":[{"name":"","start":"YYYY-MM-DD","end":"YYYY-MM-DD","owner":"","status":"Ej startad","color":"#16a34a","bar_text":"","text_position":"Automatisk"}]}\nDagens datum är %s. Beskrivning: %s'''%(date.today().isoformat(),user_text)
                else:
                    if not p: raise ValueError('Markera först ett projekt i huvudfönstret.')
                    prompt='''Föreslå en reviderad version av projektet enligt instruktionen. Returnera ENDAST hela projektet som JSON, med samma id och befintliga fält bevarade när de inte ska ändras. Ändra aldrig order_no, project_no, handler eller owner om användaren inte uttryckligen ber om det.\nINSTRUKTION:\n%s\nPROJEKT:\n%s'''%(user_text,self.project_context(p))
                proposal=extract_json(self.ai_call(prompt)); state['proposal']=proposal
                self.after(0,lambda:(setout(json.dumps(proposal,ensure_ascii=False,indent=2)),status.config(text='Förslag klart – granska innan godkännande'),apply_btn.config(state='normal'),send_btn.config(state='normal')))
            except Exception as ex: self.after(0,lambda ex=ex:(status.config(text='Fel'),send_btn.config(state='normal'),messagebox.showerror('AI-fel',str(ex),parent=win)))
        def send(event=None):
            user_text=entry.get('1.0','end').strip()
            if not user_text:return 'break'
            selected_mode=mode.get(); entry.delete('1.0','end'); state['proposal']=None; apply_btn.config(state='disabled'); send_btn.config(state='disabled'); status.config(text='AI arbetar…')
            if selected_mode=='Chatt': history.append({'role':'user','text':user_text,'at':now()}); self.save_chat(history); render_chat()
            threading.Thread(target=worker,args=(user_text,selected_mode),daemon=True).start(); return 'break'
        def apply():
            q=state.get('proposal')
            if not isinstance(q,dict): return
            if not messagebox.askyesno('Bekräfta ändring','Vill du genomföra det granskade AI-förslaget?',parent=win): return
            if mode.get()=='Skapa nytt projekt med AI': q['id']=uuid.uuid4().hex; q['revision']=0; q.setdefault('activities',[]); self.normalize(q); self.projects.append(q); self.changed(q); messagebox.showinfo('AI','Projektutkastet har lagts till.',parent=win)
            else:
                if not p:return
                keep_id=p['id']; p.clear(); p.update(q); p['id']=keep_id; self.normalize(p); self.changed(p); messagebox.showinfo('AI','Förslaget har lagts in i projektet.',parent=win)
            state['proposal']=None; apply_btn.config(state='disabled')
        apply_btn.config(command=apply); send_btn.config(command=send); entry.bind('<Control-Return>',send); entry.focus_set()

    def ai_result_ui(self, ans, proposal, setout, status, apply_btn, state):
        state['proposal']=proposal; setout(json.dumps(proposal,ensure_ascii=False,indent=2) if proposal else ans); status.config(text='Klart'); apply_btn.config(state='normal' if proposal else 'disabled')

    def sync_project_templates(self,p,quiet=False):
        """Copy templates for selected trades into the project's document area without overwriting existing files."""
        copied=[]; missing=[]; root=self.project_doc_dir(p)
        for trade in p.get('trades',[]):
            src=self.template_dir/trade
            if not src.exists(): missing.append(trade); continue
            dst=root/'04 Egenkontroller'/trade; dst.mkdir(parents=True,exist_ok=True)
            for f in src.rglob('*'):
                if not f.is_file(): continue
                rel=f.relative_to(src); target=dst/rel; target.parent.mkdir(parents=True,exist_ok=True)
                if not target.exists(): shutil.copy2(f,target); copied.append(str(target.relative_to(root)))
        # standard folders
        for name in ['01 Avtal & AF','02 Ritningar','03 Arbetsberedningar','04 Egenkontroller','05 KMA','06 ÄTA','07 Protokoll','08 Bilder','09 Övrigt']:
            (root/name).mkdir(parents=True,exist_ok=True)
        if not quiet: messagebox.showinfo('Mallimport',f'{len(copied)} nya mallfiler kopierades.' + (('\nSaknade mallmappar: '+', '.join(missing)) if missing else ''))
        return copied

    def extract_template_text(self,path):
        ext=path.suffix.lower()
        if ext in ('.txt','.md','.csv'):
            return path.read_text(encoding='utf-8',errors='ignore')[:30000]
        if ext=='.docx':
            try:
                from docx import Document
                d=Document(str(path)); parts=[x.text for x in d.paragraphs]
                for table in d.tables:
                    for row in table.rows: parts.append(' | '.join(c.text for c in row.cells))
                return '\n'.join(parts)[:30000]
            except Exception as ex: raise ValueError('Kunde inte läsa Word-mallen: '+str(ex))
        raise ValueError('AI-läsning stöder DOCX, TXT, MD och CSV. PDF-mallar kopieras till projektet men behöver konverteras till DOCX/TXT för AI-generering.')

    def generate_control_docx(self,p,template_path,ai_text):
        from docx import Document
        from docx.shared import Pt
        outdir=self.project_doc_dir(p)/'04 Egenkontroller'/'AI-genererade'; outdir.mkdir(parents=True,exist_ok=True)
        safe=re.sub(r'[^A-Za-z0-9ÅÄÖåäö _-]+','',template_path.stem).strip() or 'Egenkontroll'
        out=outdir/f'{safe}_AI_{date.today().isoformat()}.docx'
        d=Document(); title=d.add_heading('Egenkontroll',0); d.add_paragraph(f"Projekt: {p.get('name','')}")
        d.add_paragraph(f"Projektnr: {p.get('project_no','')}    Ordernr: {p.get('order_no','')}")
        d.add_paragraph(f"Handläggare: {p.get('handler','')}    Ansvarig: {p.get('owner','')}")
        d.add_paragraph(f"Underlag/mall: {template_path.name}")
        d.add_heading('Kontrollpunkter / innehåll',level=1)
        for line in ai_text.splitlines():
            line=line.strip()
            if not line: continue
            if line.startswith(('-', '•','*')): d.add_paragraph(line.lstrip('-•* ').strip(),style='List Bullet')
            else: d.add_paragraph(line)
        d.add_paragraph('\nDatum: ____________________    Signatur: ____________________')
        d.save(str(out)); return out

    def document_manager(self):
        p,_=self.selected()
        if not p: messagebox.showinfo('Välj projekt','Markera först ett projekt eller en aktivitet i projektet.'); return
        self.sync_project_templates(p,quiet=True); root=self.project_doc_dir(p)
        win=tk.Toplevel(self); win.title('Dokument – '+p['name']); win.geometry('1000x650'); win.transient(self)
        outer=ttk.Frame(win,padding=12); outer.pack(fill='both',expand=True)
        head=ttk.Frame(outer); head.pack(fill='x'); ttk.Label(head,text=p['name'],font=self.font_title).pack(side='left'); ttk.Label(head,text='Yrkesgrupper: '+(', '.join(p.get('trades',[])) or 'inga valda'),style='Sub.TLabel').pack(side='right')
        tree=ttk.Treeview(outer,columns=('type','size'),show='tree headings'); tree.heading('#0',text='Dokument'); tree.heading('type',text='Typ'); tree.heading('size',text='Storlek'); tree.column('#0',width=620); tree.column('type',width=100); tree.column('size',width=100); tree.pack(fill='both',expand=True,pady=10)
        paths={}
        def refresh_docs():
            tree.delete(*tree.get_children()); paths.clear()
            def add_dir(parent,folder):
                iid='d'+uuid.uuid4().hex; paths[iid]=folder; tree.insert(parent,'end',iid=iid,text=folder.name,values=('Mapp',''),open=True)
                for x in sorted(folder.iterdir(),key=lambda q:(q.is_file(),q.name.lower())):
                    if x.is_dir(): add_dir(iid,x)
                    else:
                        fid='f'+uuid.uuid4().hex; paths[fid]=x; tree.insert(iid,'end',iid=fid,text=x.name,values=(x.suffix.upper().lstrip('.'),f'{x.stat().st_size/1024:.0f} KB'))
            for folder in sorted([x for x in root.iterdir() if x.is_dir()],key=lambda q:q.name.lower()): add_dir('',folder)
        def selected_path():
            sel=tree.selection(); return paths.get(sel[0]) if sel else None
        def open_sel(e=None):
            x=selected_path()
            if not x:return
            open_native(x)
        tree.bind('<Double-1>',open_sel)
        bar=ttk.Frame(outer); bar.pack(fill='x')
        ttk.Button(bar,text='Öppna',command=open_sel).pack(side='left',padx=3)
        def add_file():
            src=filedialog.askopenfilename(parent=win)
            if not src:return
            target_dir=selected_path(); target_dir=target_dir if target_dir and target_dir.is_dir() else root/'09 Övrigt'; target=target_dir/Path(src).name
            if target.exists() and not messagebox.askyesno('Finns redan','Filen finns redan. Ersätta?',parent=win):return
            shutil.copy2(src,target); refresh_docs()
        ttk.Button(bar,text='Lägg till fil…',command=add_file).pack(side='left',padx=3)
        ttk.Button(bar,text='Hämta valda yrkesmallar',command=lambda:(self.sync_project_templates(p),refresh_docs())).pack(side='left',padx=3)
        def open_templates():
            open_native(self.template_dir)
        ttk.Button(bar,text='Öppna Mallar-mappen',command=open_templates).pack(side='left',padx=3)
        def ai_control():
            x=selected_path()
            if not x or not x.is_file(): messagebox.showinfo('Välj mall','Markera en DOCX/TXT/MD/CSV-mall i dokumentlistan.',parent=win); return
            try: tmpl=self.extract_template_text(x)
            except Exception as ex: messagebox.showerror('Mall',str(ex),parent=win); return
            status.config(text='AI skapar egenkontroll…')
            prompt=("Skapa ett projektspecifikt utkast till egenkontroll. Utgå strikt från mallen nedan och behåll dess kontrollpunkter och avsikt. "
                    "Du får anpassa projektuppgifter och formuleringar men ska inte hitta på myndighetskrav, standardkrav eller mätvärden som inte finns i mallen/projektinformationen. "
                    "Markera osäkra eller saknade uppgifter med [FYLL I]. Svara med svensk dokumenttext, utan markdown-tabeller.\n\nPROJEKT:\n"+self.project_context(p)+"\n\nMALL:\n"+tmpl)
            def work():
                try:
                    ans=self.ai_call(prompt); out=self.generate_control_docx(p,x,ans)
                    self.after(0,lambda:(status.config(text='Klar: '+out.name),refresh_docs(),messagebox.showinfo('Egenkontroll skapad',str(out),parent=win)))
                except Exception as ex:self.after(0,lambda ex=ex:(status.config(text='Fel'),messagebox.showerror('AI-fel',str(ex),parent=win)))
            threading.Thread(target=work,daemon=True).start()
        ttk.Button(bar,text='✨ Skapa egenkontroll från vald mall',command=ai_control).pack(side='right',padx=3)
        status=ttk.Label(outer,text=f'Dokumentmapp: {root}',style='Sub.TLabel'); status.pack(fill='x',pady=(8,0)); refresh_docs()

    def export_package(self):
        out=filedialog.asksaveasfilename(defaultextension='.zip',filetypes=[('ZIP-paket','*.zip')],initialfile=f'Ramavtalade_tidsplaner_{date.today()}.zip')
        if not out:return
        with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as z:
            for f in self.pdir.glob('*.json'):z.write(f,arcname='projects/'+f.name)
        messagebox.showinfo('Export klar','Projektpaketet har exporterats.')
    def import_package(self):
        src=filedialog.askopenfilename(filetypes=[('ZIP-paket','*.zip')]);
        if not src:return
        try:
            with zipfile.ZipFile(src) as z:
                for n in z.namelist():
                    if n.startswith('projects/') and n.endswith('.json'):
                        data=json.loads(z.read(n).decode('utf-8')); self.normalize(data); data['id']=uuid.uuid4().hex if (self.pdir/f"{data['id']}.json").exists() else data['id']; atomic_json(self.pdir/f"{data['id']}.json",data)
            self.load_all()
        except Exception as ex:messagebox.showerror('Importfel',str(ex))
    def settings(self):
        win=tk.Toplevel(self); win.title('Inställningar'); win.transient(self); win.grab_set(); f=ttk.Frame(win,padding=18); f.pack()
        u=tk.StringVar(value=self.user); ttk.Label(f,text='Ditt namn').grid(row=0,column=0,sticky='w',pady=5); ttk.Entry(f,textvariable=u,width=35).grid(row=0,column=1,pady=5)
        ttk.Label(f,text='Gemensam datamapp').grid(row=1,column=0,sticky='w',pady=5); path=ttk.Label(f,text=str(self.sync),wraplength=420); path.grid(row=1,column=1,sticky='w',pady=5)
        ttk.Checkbutton(f,text='Spara ändringar automatiskt',variable=self.autosave).grid(row=2,column=1,sticky='w',pady=8)
        ttk.Label(f,text='Företagslogga för PDF').grid(row=3,column=0,sticky='w',pady=5); logo_lbl=ttk.Label(f,text=str(self.logo) if self.logo else 'Ingen logga vald',wraplength=420); logo_lbl.grid(row=3,column=1,sticky='w',pady=5)
        def choose_logo():
            p=filedialog.askopenfilename(parent=win,filetypes=[('Bildfiler','*.png *.jpg *.jpeg')])
            if p:self.logo=Path(p); logo_lbl.config(text=str(self.logo))
        ttk.Button(f,text='Välj logga…',command=choose_logo).grid(row=4,column=1,sticky='w',pady=5)
        ttk.Separator(f,orient='horizontal').grid(row=5,column=0,columnspan=2,sticky='ew',pady=10)
        provider=tk.StringVar(value=self.cfg.get('ai_provider','Gemini')); key=tk.StringVar(value=self.cfg.get('ai_key','')); model=tk.StringVar(value=self.cfg.get('ai_model','gemini-3.5-flash'))
        ttk.Label(f,text='AI-leverantör').grid(row=6,column=0,sticky='w',pady=5); pcb=ttk.Combobox(f,textvariable=provider,values=['Gemini','OpenAI'],state='readonly',width=32); pcb.grid(row=6,column=1,sticky='w',pady=5)
        ttk.Label(f,text='API-nyckel').grid(row=7,column=0,sticky='w',pady=5); ttk.Entry(f,textvariable=key,width=38,show='•').grid(row=7,column=1,sticky='w',pady=5)
        ttk.Label(f,text='AI-modell').grid(row=8,column=0,sticky='w',pady=5); ttk.Entry(f,textvariable=model,width=38).grid(row=8,column=1,sticky='w',pady=5)
        ttk.Label(f,text='Exempel: gemini-3.5-flash eller gpt-5-mini',style='Sub.TLabel').grid(row=9,column=1,sticky='w'); ttk.Label(f,text='Mallbibliotek').grid(row=10,column=0,sticky='w',pady=5); ttk.Label(f,text=str(self.template_dir),wraplength=420).grid(row=10,column=1,sticky='w',pady=5)
        def provider_changed(e=None): model.set('gemini-3.5-flash' if provider.get()=='Gemini' else 'gpt-5-mini')
        pcb.bind('<<ComboboxSelected>>',provider_changed)
        def test_ai():
            self.cfg.update({'ai_provider':provider.get(),'ai_key':key.get().strip(),'ai_model':model.get().strip()}); self.save_cfg()
            def work():
                try: r=self.ai_call('Svara exakt med: Anslutning OK'); self.after(0,lambda:messagebox.showinfo('AI-test',r[:300],parent=win))
                except Exception as ex:self.after(0,lambda ex=ex:messagebox.showerror('AI-test',str(ex),parent=win))
            threading.Thread(target=work,daemon=True).start()
        ttk.Button(f,text='Testa AI-anslutning',command=test_ai).grid(row=11,column=1,sticky='w',pady=5)
        def choose():
            p=filedialog.askdirectory(parent=win)
            if p:self.sync=Path(p)/'RamavtaladeTidsplanerData'; self.save_cfg(); self.sync.mkdir(parents=True,exist_ok=True); (self.sync/'projects').mkdir(exist_ok=True); (self.sync/'archive').mkdir(exist_ok=True); (self.sync/'mallar').mkdir(exist_ok=True); (self.sync/'documents').mkdir(exist_ok=True); path.config(text=str(self.sync))
        ttk.Button(f,text='Byt OneDrive-mapp…',command=choose).grid(row=12,column=1,sticky='w',pady=5)
        ttk.Separator(f,orient='horizontal').grid(row=13,column=0,columnspan=2,sticky='ew',pady=10)
        lic=_license_state()
        licinfo=f"{lic.get('customer','Okänt företag')} — {lic.get('display_name','')} — {lic.get('role','')}"
        ttk.Label(f,text='Licens').grid(row=14,column=0,sticky='w',pady=5); ttk.Label(f,text=licinfo,wraplength=420).grid(row=14,column=1,sticky='w',pady=5)
        def change_license():
            if account_login(win,force_new=True):
                messagebox.showinfo('Licens','Du är inloggad med ett nytt konto. Programmet startas om för att läsa in företag, roll och OneDrive-inställningar.',parent=win)
                self.destroy()
                os.execl(sys.executable, sys.executable, *sys.argv)
        ttk.Button(f,text='Byt konto…',command=change_license).grid(row=15,column=1,sticky='w',pady=5)
        def ok(): self.user=u.get().strip() or getpass.getuser(); self.cfg.update({'ai_provider':provider.get(),'ai_key':key.get().strip(),'ai_model':model.get().strip()}); self.save_cfg(); win.destroy(); self.load_all()
        ttk.Button(f,text='Spara',command=ok).grid(row=16,column=1,sticky='e',pady=(12,0))
    def on_close(self):
        if self.dirty:
            r=messagebox.askyesnocancel('Osparade ändringar','Spara ändringarna innan programmet stängs?')
            if r is None:return
            if r:self.save_all()
        self.destroy()

if __name__=='__main__':
    gate=tk.Tk(); gate.withdraw()
    if account_login(gate):
        gate.destroy(); app=App(); app.protocol('WM_DELETE_WINDOW',app.on_close); app.mainloop()
    else:
        gate.destroy()

